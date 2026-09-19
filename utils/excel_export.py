"""Builds the .xlsx booking ledger used both for the manual admin download
and the automated daily email attachment."""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import models

HEADER_FILL = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)

COLUMNS = [
    ("Booking Date", 14), ("Meal Type", 12), ("Food Type", 10), ("Item", 22),
    ("Target", 8), ("Employee ID", 12), ("Employee Name", 20), ("Department", 18),
    ("Plant", 16), ("Guest Type", 18), ("Quantity", 10), ("Booked By", 14),
    ("Booked At", 18), ("Status", 10),
]


def _autofit(ws):
    for idx, (name, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_ledger_workbook(start_date: date, end_date: date, department=None, target=None, plant=None,
                           meal_types=None, title=None):
    """Returns an openpyxl Workbook containing every booking between
    start_date and end_date (inclusive), optionally filtered. `meal_types`,
    if given, keeps only bookings whose meal_type is in that list/set —
    applied the same way across every date in the range."""
    bookings = models.list_bookings(
        start_date=start_date, end_date=end_date, department=department, target=target, plant=plant, status=None,
    )
    if meal_types:
        allowed = set(meal_types)
        bookings = [b for b in bookings if b.meal_type in allowed]

    return _build_workbook(bookings, title or f"Food Booking Ledger ({start_date.isoformat()} to {end_date.isoformat()})")


def build_meal_slice_workbook(segments, title):
    """Builds a workbook from an explicit list of (date, meal_type) pairs,
    each fetched independently and concatenated into one ledger. This is
    what lets a single report span more than one day *and* pin a different
    meal type per day — e.g. today's Dinner together with tomorrow's
    Breakfast, since booking for both closes around the same cutoff moment.
    (build_ledger_workbook's `meal_types` filter can't express that, because
    it applies the same allowed set across the whole date range.)"""
    bookings = []
    for booking_date, meal_type in segments:
        bookings.extend(models.list_bookings(start_date=booking_date, end_date=booking_date, status=None, meal_type=meal_type))
    bookings.sort(key=lambda b: (b.booking_date, b.meal_type, b.department))
    return _build_workbook(bookings, title)


def _build_workbook(bookings, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "Booking Ledger"

    ws.merge_cells("A1:N1")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT

    header_row = 3
    for idx, (name, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=header_row, column=idx, value=name)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    row = header_row + 1
    total_self, total_guest = 0, 0
    for b in bookings:
        ws.cell(row=row, column=1, value=b.booking_date.strftime("%Y-%m-%d"))
        ws.cell(row=row, column=2, value=b.meal_type)
        ws.cell(row=row, column=3, value=b.food_type)
        ws.cell(row=row, column=4, value=b.item_name or "")
        ws.cell(row=row, column=5, value=b.target)
        ws.cell(row=row, column=6, value=b.employee.employee_code if b.employee else "")
        ws.cell(row=row, column=7, value=b.employee.name if b.employee else "")
        ws.cell(row=row, column=8, value=b.department)
        ws.cell(row=row, column=9, value=b.plant)
        ws.cell(row=row, column=10, value=b.guest_classification or "")
        ws.cell(row=row, column=11, value=b.quantity)
        ws.cell(row=row, column=12, value=b.booker.username if b.booker else "")
        ws.cell(row=row, column=13, value=b.booked_at.strftime("%Y-%m-%d %H:%M") if hasattr(b.booked_at, "strftime") else (b.booked_at or ""))
        ws.cell(row=row, column=14, value=b.status)
        if b.target == "Self":
            total_self += b.quantity
        else:
            total_guest += b.quantity
        row += 1

    _autofit(ws)

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Summary"
    ws2["A1"].font = TITLE_FONT
    ws2["A3"] = "Total meals (Self)"
    ws2["B3"] = total_self
    ws2["A4"] = "Total meals (Guest)"
    ws2["B4"] = total_guest
    ws2["A5"] = "Total meals (All)"
    ws2["B5"] = total_self + total_guest
    ws2["A6"] = "Total booking records"
    ws2["B6"] = len(bookings)
    ws2.column_dimensions["A"].width = 24

    # Department-wise breakdown
    ws2["A8"] = "Department-wise consumption"
    ws2["A8"].font = Font(bold=True)
    ws2["A9"], ws2["B9"] = "Department", "Meals"
    dept_totals = {}
    for b in bookings:
        dept_totals[b.department] = dept_totals.get(b.department, 0) + b.quantity
    r = 10
    for dept, qty in sorted(dept_totals.items(), key=lambda x: -x[1]):
        ws2.cell(row=r, column=1, value=dept)
        ws2.cell(row=r, column=2, value=qty)
        r += 1

    # Plant-wise breakdown
    r += 1
    ws2.cell(row=r, column=1, value="Plant-wise consumption").font = Font(bold=True)
    r += 1
    ws2.cell(row=r, column=1, value="Plant")
    ws2.cell(row=r, column=2, value="Meals")
    r += 1
    plant_totals = {}
    for b in bookings:
        plant_totals[b.plant] = plant_totals.get(b.plant, 0) + b.quantity
    for plant_name, qty in sorted(plant_totals.items(), key=lambda x: -x[1]):
        ws2.cell(row=r, column=1, value=plant_name)
        ws2.cell(row=r, column=2, value=qty)
        r += 1

    return wb


def workbook_to_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
